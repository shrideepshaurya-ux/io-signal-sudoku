import io
import json

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from src import excel_reader as er, report_generator as rg, revision_diff as rd, sample_data as sd
from src import allocator, validator
from src.hardware import Topology
from src.models import RuleConfig
from src.utilization import compute_utilization
from tests.conftest import run


def by_tag(res):
    return {s.tag: res.placements.get(s.sid) for s in res.signals}


# 1 --------------------------------------------------------------------------
def test_basic_allocation_complete_and_unique(cfg):
    rr, res, v = run(sd.synthetic_basic(), cfg)
    assert res.status == "OPTIMAL" and res.complete
    assert len(res.placements) == len(rr.signals)
    locs = [p.loc for p in res.placements.values()]
    assert len(locs) == len(set(locs))
    assert v.fully_valid and not v.violations


# 2 --------------------------------------------------------------------------
@pytest.mark.parametrize("hw", [{}, dict(ciocs_per_cabinet=2, baseplates_per_cioc=2, charms_per_baseplate=6)])
def test_capacity_limits(cfg, hw):
    cfg.hardware = cfg.hardware.model_copy(update=hw)
    cfg.spare.enabled = False
    df = sd.synthetic_large(400, seed=3, group_fraction=0.0) if hw else sd.synthetic_large(600, seed=3)
    if hw:   # drop cables too big for a 12-slot CIOC
        df = df[df["Cable ID"].map(df["Cable ID"].value_counts()) <= 12]
    rr, res, v = run(df, cfg)
    h = cfg.hardware
    assert v.fully_valid, [i.message for i in v.violations][:3]
    for p in res.placements.values():
        assert p.cioc < h.ciocs_per_cabinet and p.bp < h.baseplates_per_cioc and p.slot < h.charms_per_baseplate
    assert len({p.loc for p in res.placements.values()}) == len(res.placements)


# 3 --------------------------------------------------------------------------
def test_ab_segregation(cfg):
    rr, res, v = run(sd.synthetic_large(500, seed=11), cfg)
    seg = {}
    for sid, p in res.placements.items():
        s = res.signals[sid]
        if s.seg:
            seg.setdefault((p.cab, p.cioc), set()).add(s.seg)
    assert seg and all(len(x) == 1 for x in seg.values())
    assert v.fully_valid


# 4 --------------------------------------------------------------------------
@pytest.mark.parametrize("cont,sp", [("linear", "compact"), ("linear", "reserve_positions"),
                                     ("baseplate_aligned", "compact")])
def test_cable_continuity(cfg, cont, sp):
    cfg.cable.continuity, cfg.cable.spare_policy = cont, sp
    rr, res, v = run(sd.synthetic_basic(), cfg)
    assert v.fully_valid, [i.message for i in v.violations]
    topo = Topology(cfg.hardware, cfg.naming)
    cables = {}
    for s in res.signals:
        cables.setdefault(s.cable_key, []).append(s)
    for key, rows in cables.items():
        rows.sort(key=lambda r: r.pair)
        offs = [topo.offset(res.placements[r.sid]) for r in rows]
        assert len({(res.placements[r.sid].cab, res.placements[r.sid].cioc) for r in rows}) == 1
        assert offs == sorted(offs)
        if sp == "compact":
            assert offs[-1] - offs[0] == len(rows) - 1
        else:
            assert [o - offs[0] for o in offs] == [r.pair - rows[0].pair for r in rows]


def test_spare_pairs_reported_not_signals(cfg):
    rr, res, v = run(sd.synthetic_basic(), cfg)
    t = {r["Cable"]: r for r in res.cable_blocks}
    assert t["CAB-001"]["Cable size (pairs)"] == 24 and t["CAB-001"]["Spare pairs"] == 2
    assert t["CAB-005"]["Spare pairs"] == 0


# 5 --------------------------------------------------------------------------
def test_unknown_io_and_missing_compat_not_assigned(cfg):
    rr, res, v = run(sd.synthetic_with_problems(), cfg)
    assert by_tag(res)["UNK-01"] is None
    assert any(t.category == "Unknown IO type" for t in res.tqs)
    del cfg.compat["DO|RELAY"]
    rr, res, v = run(sd.synthetic_basic(), cfg)
    assert all(by_tag(res)[f"RLY601-{i:02d}"] is None for i in range(1, 7))
    assert any(t.category == "Unknown CHARM compatibility" for t in res.tqs)


def test_charm_mapping_follows_matrix(cfg):
    rr, res, v = run(sd.synthetic_basic(), cfg)
    p = by_tag(res)
    assert p["SOL501-01"].charm_type == "DO ISOLATED CHARM"
    assert p["SOL501-01"].terminal_block == "FUSED INJECTED POWER TERMINAL BLOCK"
    assert p["RLY601-01"].terminal_block == "RELAY OUTPUT TERMINAL BLOCK"
    assert p["PT101-01"].charm_type == "AI CHARM"
    assert p["FT701-01"].configuration == "HART"


# 6 --------------------------------------------------------------------------
def test_missing_mandatory_fields(cfg):
    rr, res, v = run(sd.synthetic_with_problems(), cfg)
    rep = rr.report
    assert rep.missing_required >= 2 and rep.missing_cable_or_pair >= 2
    p = by_tag(res)
    assert p["NOCABLE-01"] is None and p["NOPAIR-01"] is None
    assert rep.total_rows == len(rr.signals) == len(res.signals)   # nothing silently dropped


# 7 --------------------------------------------------------------------------
def test_duplicate_tags_flagged_not_overwritten(cfg):
    rr, res, v = run(sd.synthetic_with_problems(), cfg)
    dups = [s for s in res.signals if s.tag == "PT101-01"]
    assert len(dups) == 2 and all(s.sid not in res.placements for s in dups)
    assert rr.report.duplicate_tags == 2
    cfg.duplicate_policy = "keep_first"
    rr, res, v = run(sd.synthetic_with_problems(), cfg)
    dups = [s for s in res.signals if s.tag == "PT101-01"]
    assert sum(s.sid in res.placements for s in dups) == 1


# 8 --------------------------------------------------------------------------
def test_infeasible_capacity_gives_status_and_tqs(cfg):
    cfg.hardware.max_cabinets = 1
    cfg.spare.enabled = False
    cfg.grouping.enabled = False
    rr, res, v = run(sd.synthetic_large(700, seed=5, group_fraction=0), cfg)
    assert res.status == "INFEASIBLE" and not res.complete
    assert 0 < len(res.placements) < len(res.signals)
    assert not v.fully_valid                       # incomplete must never be 'fully valid'
    assert not v.violations                        # but what was placed is itself valid
    cap = [t for t in res.tqs if t.category == "Insufficient hardware capacity"]
    assert cap and "available" in cap[0].explanation
    assert all(s.sid in res.placements or s.sid in res.unallocated for s in res.signals)


def test_ab_cable_conflict_tq(cfg):
    rr, res, v = run(sd.synthetic_with_problems(), cfg)
    t = [t for t in res.tqs if t.category == "Conflicting A/B segregation"]
    assert t and "CAB-105" in t[0].explanation
    assert by_tag(res)["MIX-01"] is None


def test_group_too_large_tq(cfg):
    cfg.hardware.max_cabinets = 2
    rows = []
    for i in range(30):
        rows += sd._rows_for_cable(f"CG{i:02d}", 12, 12, "DI", group="GBIG", prefix=f"G{i:02d}")
    rr, res, v = run(pd.DataFrame(rows, columns=sd.COLUMNS), cfg)
    assert any(t.category == "Insufficient hardware capacity" and "GBIG" in t.subject for t in res.tqs)
    assert not res.placements


# 9 --------------------------------------------------------------------------
def test_spare_capacity_empty_baseplates(cfg):
    cfg.spare.fraction = 0.25            # 8 of 32 baseplates
    rr, res, v = run(sd.synthetic_large(1200, seed=2), cfg)
    assert v.fully_valid
    u = compute_utilization(res, cfg)
    assert (u.cabinets["Empty baseplates"] >= 8).all() and u.cabinets["Spare met"].all()


def test_spare_shortfall_reported_when_cabinets_limited(cfg):
    cfg.grouping.enabled = False
    cfg.spare.fraction = 0.5
    cfg.hardware.max_cabinets = 2
    rr, res, v = run(sd.synthetic_large(500, seed=9, group_fraction=0), cfg)
    assert res.complete
    assert res.spare["shortfall_baseplates_per_cabinet"] > 0
    assert any(f.code == "SPARE_SHORTFALL" for f in res.findings)
    assert any(t.category == "Spare-capacity shortfall" for t in res.tqs)


# 10 -------------------------------------------------------------------------
def test_revision_preserves_unchanged_locations(cfg):
    rr, res_a, _ = run(sd.synthetic_basic(), cfg)
    rb = er.read_master(sd.to_workbook_bytes(sd.synthetic_revision_b()), "b.xlsx", cfg)
    rep = rd.compare_revisions(res_a, rb.signals, cfg, "B")
    s = rep.summary
    assert s["Added"] == 10 and s["Deleted"] == 2 and s["Modified"] == 2
    assert s["Reallocated"] == 0
    pa, pb = by_tag(res_a), by_tag(rep.result_b)
    unchanged = rep.changes[rep.changes["Change"] == "Unchanged signal"]["Tag"]
    assert len(unchanged) > 50
    assert all(pa[t].loc == pb[t].loc for t in unchanged)
    assert validator.validate(rep.result_b, cfg).fully_valid


def test_revision_restore_from_workbook(cfg):
    rr, res_a, v = run(sd.synthetic_basic(), cfg)
    data, _ = rg.export_workbook(res_a, cfg, v, rr.report)
    prev = rd.previous_from_workbook(data, cfg)
    assert len(prev) == len(res_a.placements)
    t = res_a.signals[0]
    assert prev[t.tag_key].loc == res_a.placements[t.sid].loc


# 11 -------------------------------------------------------------------------
def test_excel_export_provisional_and_template(cfg):
    rr, res, v = run(sd.synthetic_with_problems(), cfg)
    data, info = rg.export_workbook(res, cfg, v, rr.report, source_files=["x.xlsx"])
    wb = load_workbook(io.BytesIO(data))
    for name in ["Allocation", "Input Validation", "Rule Checks", "TQs", "Cabinet Utilization",
                 "CIOC Utilization", "Baseplate Utilization", "Cable Spares", "Solver Summary", "Audit"]:
        assert name in wb.sheetnames, name
    assert info["provisional"]
    alloc = pd.DataFrame(wb["Allocation"].values)
    assert len(alloc) >= len(res.signals) + 2         # note + header + rows (nothing dropped)
    # official-style template with custom order, extra column, merged title
    t = Workbook()
    ws = t.active
    ws.title = "IO List"
    ws["A1"] = "PROJECT X IO ALLOCATION"
    ws.merge_cells("A1:F1")
    ws.append([])
    ws.append(["Tag", "Slot", "Cabinet", "CIOC", "Baseplate", "Owner Notes", "CHARM Type"])
    ws.append(["EXAMPLE", None, None, None, None, "keep me?", None])
    buf = io.BytesIO()
    t.save(buf)
    data, info = rg.export_workbook(res, cfg, v, rr.report, template=buf.getvalue())
    assert info["template_used"] and info["template_sheet"] == "IO List"
    ws = load_workbook(io.BytesIO(data))["IO List"]
    assert ws["A1"].value == "PROJECT X IO ALLOCATION"
    assert "A1:F1" in [str(r) for r in ws.merged_cells.ranges]
    hdr = next(r for r in range(1, 10) if ws.cell(row=r, column=1).value == "Tag")
    assert ws.cell(row=hdr + 1, column=1).value == res.signals[0].tag
    assert "Allocation Detail" in load_workbook(io.BytesIO(data)).sheetnames


def test_formula_injection_is_text(cfg):
    df = sd.synthetic_basic()
    df.loc[0, "Signal Tag"] = '=HYPERLINK("x")'
    rr, res, v = run(df, cfg)
    data, _ = rg.export_workbook(res, cfg, v, rr.report)
    ws = load_workbook(io.BytesIO(data))["Allocation"]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)


# 12 -------------------------------------------------------------------------
def test_validator_catches_duplicate_slot(cfg):
    rr, res, v = run(sd.synthetic_basic(), cfg)
    assert v.fully_valid
    t = by_tag(res)
    sid = {s.tag: s.sid for s in res.signals}
    p = res.placements[sid["PT101-02"]]
    p.cab, p.cioc, p.bp, p.slot = t["PT101-01"].cab, t["PT101-01"].cioc, t["PT101-01"].bp, t["PT101-01"].slot
    v = validator.validate(res, cfg)
    assert not v.fully_valid and any(i.rule_id == "R-CAP" for i in v.violations)


def test_validator_catches_segregation_cable_charm_and_missing(cfg):
    rr, res, _ = run(sd.synthetic_basic(), cfg)
    sid = {s.tag: s.sid for s in res.signals}
    a = res.placements[sid["PT101-01"]]
    b = res.placements[sid["ZS201-01"]]
    b.cab, b.cioc, b.bp, b.slot = a.cab, a.cioc, 7, 11
    v = validator.validate(res, cfg)
    ids = {i.rule_id for i in v.violations}
    assert "R-SEG" in ids and "R-CABLE" in ids
    res.placements[sid["PT101-03"]].charm_type = "WRONG CHARM"
    assert "R-CHARM" in {i.rule_id for i in validator.validate(res, cfg).violations}
    del res.placements[sid["XV401-01"]]
    v = validator.validate(res, cfg)
    assert any(i.rule_id == "STRUCT" for i in v.violations) and not v.fully_valid


# 13 -------------------------------------------------------------------------
def test_disabling_rule_changes_validation(cfg):
    rr, res, _ = run(sd.synthetic_basic(), cfg)
    sid = {s.tag: s.sid for s in res.signals}
    a, b = res.placements[sid["PT101-01"]], res.placements[sid["ZS201-01"]]
    b.cab, b.cioc, b.bp, b.slot = a.cab, a.cioc, 7, 11
    assert any(i.rule_id == "R-SEG" for i in validator.validate(res, cfg).violations)
    cfg2 = cfg.model_copy(deep=True)
    cfg2.segregation.enabled = False
    v2 = validator.validate(res, cfg2)
    assert not any(i.rule_id == "R-SEG" for i in v2.violations)
    assert v2.rules_checked["R-SEG"] is False


def test_profile_roundtrip(tmp_path, cfg):
    from src.models import load_config, save_config
    cfg.spare.fraction = 0.2
    p = tmp_path / "p.json"
    save_config(cfg, str(p))
    assert load_config(str(p)).spare.fraction == 0.2
    assert load_config(str(p)).config_id == cfg.config_id


def test_invalid_config_rejected():
    with pytest.raises(Exception):
        RuleConfig.model_validate({"naming": {"cioc": "CIOC{nope}"}})
    with pytest.raises(Exception):
        RuleConfig.model_validate({"hardware": {"charms_per_baseplate": 0}})


# reader ------------------------------------------------------------------------
def test_reader_errors_and_manual_mapping(cfg):
    with pytest.raises(er.InputError):
        er.read_master(b"", "x.xlsx", cfg)
    with pytest.raises(er.InputError):
        er.read_master(b"not an excel", "x.xlsx", cfg)
    with pytest.raises(er.InputError):
        er.read_master(b"abc", "x.csv", cfg)
    wb = Workbook()
    wb.active.append(["Instrument", "Type", "Cbl", "Core"])
    wb.active.append(["T1", "AI", "C1", 1])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(er.InputError):   # unfamiliar headers need manual mapping
        er.read_master(buf.getvalue(), "x.xlsx", cfg)
    df, hr = er.load_table(buf.getvalue(), "x.xlsx", "Sheet")
    mapping = {"tag": "Instrument", "io_type": "Type", "cable_id": "Cbl", "pair": "Core"}
    sigs, rep = er.build_signals(df, mapping, cfg)
    assert rep.valid_signals == 1 and sigs[0].tag == "T1"


def test_original_values_preserved(cfg):
    df = sd.synthetic_basic()
    df.loc[0, "Signal Tag"] = "  pt101-01  "
    df.loc[0, "IO Type"] = "ai"
    rr = er.read_master(sd.to_workbook_bytes(df), "t.xlsx", cfg)
    assert rr.signals[0].tag == "pt101-01"
    assert rr.signals[0].io_key == "AI"
    assert rr.raw_df.iloc[0]["Signal Tag"] == "  pt101-01  "


# 14 -------------------------------------------------------------------------
def test_mid_size_dataset(cfg):
    cfg.solver.time_limit_s = 20
    rr, res, v = run(sd.synthetic_large(2500, seed=1), cfg)
    assert res.complete and v.fully_valid
    assert res.solver["runtime_s"] < 60


@pytest.mark.benchmark
def test_benchmark_10000(cfg):
    import os
    import time
    import psutil
    cfg.solver.time_limit_s = 60
    proc = psutil.Process(os.getpid())
    t0 = time.time()
    df = sd.synthetic_large(10000, seed=7)
    rr = er.read_master(sd.to_workbook_bytes(df), "big.xlsx", cfg)
    t_read = time.time() - t0
    t1 = time.time()
    res = allocator.allocate(rr.signals, cfg, revision="bench")
    t_alloc = time.time() - t1
    t2 = time.time()
    v = validator.validate(res, cfg)
    t_val = time.time() - t2
    out = {"signals": 10000, "read_s": round(t_read, 2), "allocate_s": round(t_alloc, 2),
           "validate_s": round(t_val, 2), "status": res.status, "complete": res.complete,
           "rss_mb": round(proc.memory_info().rss / 1e6), "solver": res.solver,
           "valid": v.fully_valid, "violations": len(v.violations), "warnings": len(v.warnings),
           "cabinets": compute_utilization(res, cfg).totals["Cabinets required"]}
    os.makedirs("outputs", exist_ok=True)
    with open("outputs/benchmark_10000.json", "w") as fh:
        json.dump(out, fh, indent=2, default=str)
    assert res.complete and v.fully_valid
