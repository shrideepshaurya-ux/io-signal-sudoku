"""Revision comparison: match by tag, classify changes, re-allocate with
movement minimisation, and explain moves."""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Optional

import pandas as pd

from .allocator import allocate
from .hardware import Topology
from .models import AllocationResult, Placement, RuleConfig, Signal
from .tq_generator import generate_tqs

ATTRS = [("io_raw", "IO type"), ("subtype_raw", "Characteristic"), ("cable_id", "Cable"),
         ("pair", "Pair"), ("group", "Group"), ("seg", "Segregation"), ("cable_pairs", "Cable size")]


@dataclass
class RevisionReport:
    summary: dict
    changes: pd.DataFrame
    moves: pd.DataFrame
    affected: pd.DataFrame
    tq_changes: pd.DataFrame
    result_b: Optional[AllocationResult] = None


def previous_from_result(res: AllocationResult) -> dict[str, Placement]:
    """tag_key -> Placement for every allocated signal of a result."""
    return {res.signals[sid].tag_key: p for sid, p in res.placements.items()
            if res.signals[sid].tag}


def previous_from_workbook(data: bytes, cfg: RuleConfig) -> dict[str, Placement]:
    """Rebuild previous placements from an exported allocation workbook
    (sheet 'Allocation' / 'Allocation Detail') by reversing the naming patterns."""
    import io
    from openpyxl import load_workbook
    topo = Topology(cfg.hardware, cfg.naming)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = None
    for name in ("Allocation Detail", "Allocation"):
        if name in wb.sheetnames:
            ws = wb[name]
            break
    if ws is None:
        raise ValueError("No 'Allocation' sheet found in the workbook.")
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and "Signal Tag" in [str(c) for c in r])
    hdr = [str(c) for c in rows[hdr_i]]
    need = ["Signal Tag", "Cabinet", "CIOC", "Baseplate", "CHARM Slot"]
    if any(n not in hdr for n in need):
        raise ValueError("Allocation sheet lacks the columns needed to restore locations.")
    ix = {n: hdr.index(n) for n in need}
    lookup_bp = {}
    lookup_cioc = {}
    lookup_cab = {}
    max_cab = 64
    for c in range(max_cab):
        lookup_cab[topo.system_cabinet(c)] = c
        for i in range(cfg.hardware.ciocs_per_cabinet):
            lookup_cioc[topo.cioc_name(c, i)] = (c, i)
            for b in range(cfg.hardware.baseplates_per_cioc):
                lookup_bp[topo.baseplate_name(c, i, b)] = (c, i, b)
    out: dict[str, Placement] = {}
    for r in rows[hdr_i + 1:]:
        if not r or r[ix["Signal Tag"]] is None:
            continue
        bp = lookup_bp.get(str(r[ix["Baseplate"]]))
        slot = r[ix["CHARM Slot"]]
        if bp is None or not isinstance(slot, (int, float)):
            continue
        out[str(r[ix["Signal Tag"]]).strip().upper()] = Placement(cab=bp[0], cioc=bp[1], bp=bp[2],
                                                                    slot=int(slot) - 1)
    return out


def _attrs(s: Signal) -> dict:
    return {k: (str(getattr(s, k)).upper() if getattr(s, k) is not None else "") for k, _ in ATTRS}


def compare_revisions(res_a: AllocationResult, signals_b: list[Signal], cfg: RuleConfig,
                      rev_b_label: str = "Rev B", previous: Optional[dict[str, Placement]] = None,
                      time_limit: Optional[float] = None, result_b: Optional[AllocationResult] = None
                      ) -> RevisionReport:
    topo = Topology(cfg.hardware, cfg.naming)
    prev = previous if previous is not None else previous_from_result(res_a)
    if result_b is None:
        result_b = allocate(signals_b, cfg, previous=prev, revision=rev_b_label, time_limit=time_limit)
        generate_tqs(result_b)
    return diff_results(res_a, result_b, cfg, prev)


def diff_results(res_a: AllocationResult, res_b: AllocationResult, cfg: RuleConfig,
                 prev: Optional[dict[str, Placement]] = None) -> RevisionReport:
    topo = Topology(cfg.hardware, cfg.naming)
    prev = prev if prev is not None else previous_from_result(res_a)

    def index(res):
        """tag_key -> list of signals (duplicates detected)."""
        d: dict[str, list[Signal]] = {}
        for s in res.signals:
            if s.tag:
                d.setdefault(s.tag_key, []).append(s)
        return d

    ia, ib = index(res_a), index(res_b)
    pb_a = res_a.placements
    pb_b = res_b.placements
    a_valid_sid = {s.sid for s in res_a.signals if not s.has_error}
    rows, moves = [], []
    # lookups for reasons -----------------------------------------------------
    cab_mem_a: dict[str, set[str]] = {}
    cab_mem_b: dict[str, set[str]] = {}
    grp_mem_a: dict[str, set[str]] = {}
    grp_mem_b: dict[str, set[str]] = {}
    for res, cm, gm in ((res_a, cab_mem_a, grp_mem_a), (res_b, cab_mem_b, grp_mem_b)):
        for s in res.signals:
            if s.cable_key:
                cm.setdefault(s.cable_key, set()).add(s.tag_key)
            if s.group:
                gm.setdefault(s.group, set()).add(s.tag_key)

    for key in sorted(set(ia) | set(ib)):
        la, lb = ia.get(key), ib.get(key)
        if la and len(la) > 1 or lb and len(lb) > 1:
            rows.append({"Tag": (la or lb)[0].tag, "Change": "Ambiguous (duplicate tag)",
                         "Details": "Tag duplicated in " + ("Rev A" if la and len(la) > 1 else "Rev B"),
                         "Previous location": "", "New location": "", "Reason": ""})
            continue
        a = la[0] if la else None
        b = lb[0] if lb else None
        if a is None:
            rows.append({"Tag": b.tag, "Change": "Added signal", "Details": "",
                         "Previous location": "",
                         "New location": topo.location_str(pb_b[b.sid]) if b.sid in pb_b else "unallocated",
                         "Reason": ""})
            continue
        if b is None:
            rows.append({"Tag": a.tag, "Change": "Deleted signal", "Details": "",
                         "Previous location": topo.location_str(pb_a[a.sid]) if a.sid in pb_a else "unallocated",
                         "New location": "", "Reason": "Removed from revised MDB"})
            continue
        aa, bb = _attrs(a), _attrs(b)
        diffs = [f"{lab}: {getattr(a, k) or '-'} -> {getattr(b, k) or '-'}"
                 for k, lab in ATTRS if aa[k] != bb[k]]
        old = pb_a.get(a.sid)
        new = pb_b.get(b.sid)
        oldl = topo.location_str(old) if old else "unallocated"
        newl = topo.location_str(new) if new else "unallocated"
        if a.sid in a_valid_sid and b.has_error:
            cls = "Newly invalid signal"
        elif diffs:
            cls = "Modified signal"
        else:
            cls = "Unchanged signal"
        reason = ""
        if old and new and old.loc != new.loc:
            if cls == "Unchanged signal":
                cls = "Reallocated signal"
            else:
                cls += " + reallocated"
            reason = _move_reason(a, b, diffs, cab_mem_a, cab_mem_b, grp_mem_a, grp_mem_b)
            moves.append({"Tag": a.tag, "Previous location": oldl, "New location": newl,
                          "Previous cabinet": topo.system_cabinet(old.cab),
                          "New cabinet": topo.system_cabinet(new.cab),
                          "Previous CIOC": topo.cioc_name(old.cab, old.cioc),
                          "New CIOC": topo.cioc_name(new.cab, new.cioc), "Reason": reason})
        rows.append({"Tag": a.tag, "Change": cls, "Details": "; ".join(diffs),
                     "Previous location": oldl, "New location": newl, "Reason": reason})
    changes = pd.DataFrame(rows, columns=["Tag", "Change", "Details", "Previous location",
                                          "New location", "Reason"])
    moves_df = pd.DataFrame(moves, columns=["Tag", "Previous location", "New location", "Previous cabinet",
                                            "New cabinet", "Previous CIOC", "New CIOC", "Reason"])
    # affected hardware ------------------------------------------------------------
    aff = set()
    for m in moves:
        aff |= {("Cabinet", m["Previous cabinet"]), ("Cabinet", m["New cabinet"]),
                ("CIOC", m["Previous CIOC"]), ("CIOC", m["New CIOC"])}
    for r in rows:
        if r["Change"] == "Added signal" and r["New location"] not in ("", "unallocated"):
            bp = r["New location"].split("/")[0]
            aff.add(("Baseplate", bp))
    affected = pd.DataFrame(sorted(aff), columns=["Level", "Name"])
    # TQ changes -------------------------------------------------------------------
    ka = {t.key: t for t in res_a.tqs}
    kb = {t.key: t for t in res_b.tqs}
    tq_rows = []
    for k, t in kb.items():
        tq_rows.append({"TQ ID (B)": t.tq_id, "Subject": t.subject, "Category": t.category,
                        "Change": "Persistent TQ" if k in ka else "Newly created TQ",
                        "Explanation": t.explanation})
    for k, t in ka.items():
        if k not in kb:
            tq_rows.append({"TQ ID (B)": "", "Subject": t.subject, "Category": t.category,
                            "Change": "Resolved TQ", "Explanation": f"Was {t.tq_id}: {t.explanation}"})
    tq_df = pd.DataFrame(tq_rows, columns=["TQ ID (B)", "Subject", "Category", "Change", "Explanation"])

    def cnt(label):
        return int(changes["Change"].str.contains(label, regex=False).sum()) if not changes.empty else 0
    unchanged_kept = int(((changes["Change"] == "Unchanged signal")).sum()) if not changes.empty else 0
    summary = {
        "Revision A": res_a.revision, "Revision B": res_b.revision,
        "Signals in Rev A": len(res_a.signals), "Signals in Rev B": len(res_b.signals),
        "Added": cnt("Added"), "Deleted": cnt("Deleted"), "Modified": cnt("Modified"),
        "Unchanged (same location)": unchanged_kept, "Reallocated": cnt("eallocated"),
        "Newly invalid": cnt("Newly invalid"),
        "New TQs": int((tq_df["Change"] == "Newly created TQ").sum()) if not tq_df.empty else 0,
        "Resolved TQs": int((tq_df["Change"] == "Resolved TQ").sum()) if not tq_df.empty else 0,
        "Persistent TQs": int((tq_df["Change"] == "Persistent TQ").sum()) if not tq_df.empty else 0,
        "Cabinets/CIOCs affected": len(affected),
        "Solver status (Rev B)": res_b.status,
    }
    return RevisionReport(summary, changes, moves_df, affected, tq_df, res_b)


def _move_reason(a: Signal, b: Signal, diffs, cm_a, cm_b, gm_a, gm_b) -> str:
    if diffs:
        return "Own attributes changed (" + "; ".join(diffs) + ")."
    if a.cable_key and cm_a.get(a.cable_key) != cm_b.get(a.cable_key):
        return "Cable composition changed (pairs added/removed), block re-laid to stay contiguous."
    if a.group and gm_a.get(a.group) != gm_b.get(a.group):
        return "Logical group membership changed; group re-packed to stay in one cabinet."
    return ("Relocated by the optimiser to satisfy hard constraints (capacity / segregation / grouping) "
            "after other changes; individual cause not determinable.")
