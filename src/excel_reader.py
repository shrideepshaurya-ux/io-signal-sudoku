"""IO MDB ingestion: worksheet selection, column detection/mapping,
normalisation and input validation. Original data is kept untouched.

Workbooks are read with ``data_only=True`` / read-only mode: formulas are never
evaluated and macros are never executed.
"""
from __future__ import annotations

import hashlib
import io
import logging
import re
from dataclasses import dataclass, field
from typing import BinaryIO, Optional, Union

import pandas as pd
from openpyxl import load_workbook

from .cable_model import derive_cable_sizes
from .models import Finding, RuleConfig, Signal

log = logging.getLogger(__name__)

ALLOWED_EXT = (".xlsx", ".xlsm")
FIELDS = ["tag", "io_type", "cable_id", "pair", "group", "seg", "subtype", "cable_pairs"]
MANDATORY_DEFAULT = ["tag", "io_type", "cable_id", "pair"]

HEADER_ALIASES: dict[str, list[str]] = {
    "tag": ["tag", "signal tag", "signal", "instrument tag", "io tag", "tag no", "tag number",
            "tagname", "tag name", "signal name"],
    "io_type": ["io type", "iotype", "signal type", "io", "type", "io signal type", "channel type"],
    "cable_id": ["cable", "cable id", "cable no", "cable number", "cable name", "cable tag", "cableid"],
    "pair": ["pair", "pair no", "pair number", "pair id", "core", "pair #", "cable pair", "pairno"],
    "group": ["group", "logical group", "io group", "grouping", "logical grouping", "group id"],
    "seg": ["segregation", "seg", "segregation group", "a b", "a/b", "redundancy",
            "segregation designation", "segregation class"],
    "subtype": ["characteristic", "signal characteristic", "subtype", "sub type", "signal class",
                "service", "io subtype", "load type"],
    "cable_pairs": ["cable pairs", "cable size", "pairs in cable", "no of pairs", "number of pairs",
                    "cable pair count"],
}


class InputError(Exception):
    """User-facing input problem (message is safe to display)."""


@dataclass
class ValidationReport:
    total_rows: int = 0
    valid_signals: int = 0
    invalid_records: int = 0
    duplicate_tags: int = 0
    missing_required: int = 0
    unrecognized_io: int = 0
    missing_cable_or_pair: int = 0
    warnings: int = 0
    issues: pd.DataFrame = field(default_factory=pd.DataFrame)

    def summary_rows(self) -> list[tuple[str, int]]:
        return [("Total input rows", self.total_rows), ("Valid signals", self.valid_signals),
                ("Invalid records", self.invalid_records),
                ("Duplicate tags (rows)", self.duplicate_tags),
                ("Rows with missing required fields", self.missing_required),
                ("Unrecognized IO types", self.unrecognized_io),
                ("Missing cable or pair information", self.missing_cable_or_pair),
                ("Warnings requiring engineering review", self.warnings)]


@dataclass
class ReadResult:
    filename: str
    sha256: str
    sheet: str
    header_row: int                     # 0-based index into sheet rows
    mapping: dict[str, Optional[str]]
    raw_df: pd.DataFrame                # untouched copy for audit
    signals: list[Signal]
    report: ValidationReport
    revision_label: str = ""

    def findings(self) -> list[Finding]:
        return [f for s in self.signals for f in s.findings]


# ---------------------------------------------------------------------------
def _norm_header(h) -> str:
    return re.sub(r"[^a-z0-9#/]+", " ", str(h).lower()).strip()


def _clean(v) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def _to_bytes(f: Union[str, bytes, BinaryIO]) -> bytes:
    if isinstance(f, (bytes, bytearray)):
        return bytes(f)
    if isinstance(f, str):
        with open(f, "rb") as fh:
            return fh.read()
    f.seek(0)
    return f.read()


def _open(data: bytes, filename: str, max_mb: int):
    if not filename.lower().endswith(ALLOWED_EXT):
        raise InputError(f"Unsupported file type: {filename!r}. Use .xlsx or .xlsm.")
    if len(data) == 0:
        raise InputError("The uploaded file is empty.")
    if len(data) > max_mb * 1024 * 1024:
        raise InputError(f"File exceeds the {max_mb} MB upload limit.")
    try:
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:  # corrupt / not a zip / etc.
        log.exception("failed to open workbook %s", filename)
        raise InputError("The file could not be opened as an Excel workbook "
                         "(it may be corrupt or password-protected).") from exc


def list_sheets(f, filename: str, max_mb: int = 50) -> list[str]:
    wb = _open(_to_bytes(f), filename, max_mb)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _sheet_rows(f, filename: str, sheet: str, max_mb: int = 50) -> list[list]:
    wb = _open(_to_bytes(f), filename, max_mb)
    try:
        if sheet not in wb.sheetnames:
            raise InputError(f"Worksheet {sheet!r} not found.")
        ws = wb[sheet]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    # strip trailing fully-empty rows
    while rows and all(c is None or str(c).strip() == "" for c in rows[-1]):
        rows.pop()
    if not rows:
        raise InputError(f"Worksheet {sheet!r} is empty.")
    return rows


def preview(f, filename: str, sheet: str, nrows: int = 25, max_mb: int = 50) -> pd.DataFrame:
    rows = _sheet_rows(f, filename, sheet, max_mb)[:nrows]
    width = max(len(r) for r in rows)
    rows = [r + [None] * (width - len(r)) for r in rows]
    return pd.DataFrame(rows).astype(object)


def detect_header_row(rows: list[list], scan: int = 30) -> int:
    all_alias = {a for al in HEADER_ALIASES.values() for a in al}
    best, best_score = 0, -1
    for i, r in enumerate(rows[:scan]):
        score = sum(1 for c in r if c is not None and _norm_header(c) in all_alias)
        if score > best_score:
            best, best_score = i, score
    return best if best_score >= 2 else 0


def detect_mapping(headers: list[str]) -> tuple[dict[str, Optional[str]], dict[str, str]]:
    """Return (mapping, confidence). confidence in {'exact','partial','none','ambiguous'}."""
    norm = {h: _norm_header(h) for h in headers if h}
    mapping: dict[str, Optional[str]] = {}
    conf: dict[str, str] = {}
    used: set[str] = set()
    for fld in FIELDS:
        aliases = HEADER_ALIASES[fld]
        exact = [h for h, n in norm.items() if n in aliases and h not in used]
        if len(exact) == 1:
            mapping[fld], conf[fld] = exact[0], "exact"
            used.add(exact[0])
            continue
        if len(exact) > 1:
            mapping[fld], conf[fld] = exact[0], "ambiguous"
            continue
        part = [h for h, n in norm.items()
                if h not in used and any(len(a) > 3 and a in n for a in aliases)]
        if len(part) == 1:
            mapping[fld], conf[fld] = part[0], "partial"
            used.add(part[0])
        else:
            mapping[fld], conf[fld] = (part[0] if part else None), ("ambiguous" if part else "none")
    return mapping, conf


def load_table(f, filename: str, sheet: str, header_row: Optional[int] = None,
               max_mb: int = 50) -> tuple[pd.DataFrame, int]:
    """Return the raw table (all cells as given) and the header row used."""
    rows = _sheet_rows(f, filename, sheet, max_mb)
    hr = detect_header_row(rows) if header_row is None else header_row
    if hr >= len(rows):
        raise InputError("Header row is beyond the end of the worksheet.")
    header = []
    seen: dict[str, int] = {}
    for i, h in enumerate(rows[hr]):
        name = _clean(h) or f"Column {i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        header.append(name)
    body = [r + [None] * (len(header) - len(r)) for r in rows[hr + 1:]]
    body = [r[:len(header)] for r in body]
    df = pd.DataFrame(body, columns=header).astype(object)
    df.insert(0, "_source_row", list(range(hr + 2, hr + 2 + len(df))))  # Excel row numbers
    keep = ~df.drop(columns="_source_row").apply(
        lambda r: all(_clean(v) is None for v in r), axis=1)
    df = df[keep].reset_index(drop=True)
    if df.empty:
        raise InputError("No data rows were found below the header row.")
    return df, hr


# ---------------------------------------------------------------------------
def resolve_io_key(io_raw: Optional[str], sub_raw: Optional[str], cfg: RuleConfig) -> Optional[str]:
    """Map raw IO type (+ optional characteristic) to a compatibility key.
    Returns None when unrecognised - never guesses."""
    if not io_raw:
        return None
    io = re.sub(r"\s+", " ", io_raw.upper().strip())
    sub = re.sub(r"[\s-]+", "_", sub_raw.upper().strip()) if sub_raw else None
    aliases = cfg.io_type_aliases
    base = aliases.get(io)
    if base is None and io in cfg.compat:
        base = io
    if base is None:
        return None
    if sub and "|" not in base:
        cand = f"{base}|{sub}"
        if cand in cfg.compat:
            return cand
        sub_alias = aliases.get(f"{io} {sub.replace('_', ' ')}")
        if sub_alias:
            return sub_alias
        return cand  # unknown combination -> reported as UNKNOWN_COMPAT downstream
    return base


def _parse_pair(v) -> tuple[Optional[int], Optional[str]]:
    s = _clean(v)
    if s is None:
        return None, None
    m = re.fullmatch(r"(?:P|PAIR\s*)?0*(\d+)(?:\.0+)?", s.upper().replace("#", ""))
    if not m:
        return None, f"Pair value {s!r} is not a positive integer"
    n = int(m.group(1))
    if n < 1:
        return None, f"Pair value {s!r} must be >= 1"
    return n, None


def build_signals(df: pd.DataFrame, mapping: dict[str, Optional[str]], cfg: RuleConfig,
                  mandatory: Optional[list[str]] = None, source_file: str = "") -> tuple[list[Signal], ValidationReport]:
    return build_signals_multi([(df, mapping, source_file)], cfg, mandatory)


def build_signals_multi(frames: list[tuple[pd.DataFrame, dict[str, Optional[str]], str]], cfg: RuleConfig,
                        mandatory: Optional[list[str]] = None) -> tuple[list[Signal], ValidationReport]:
    """Build signals from one or more mapped tables; cross-record checks
    (duplicates etc.) run once over the combined set."""
    signals: list[Signal] = []
    for df, mapping, name in frames:
        _signals_from_frame(df, mapping, cfg, mandatory, name, signals)
    _cross_checks(signals, cfg)
    return signals, _report(signals, None)


def _signals_from_frame(df: pd.DataFrame, mapping: dict[str, Optional[str]], cfg: RuleConfig,
                        mandatory: Optional[list[str]], source_file: str,
                        signals: list[Signal]) -> None:
    mandatory = list(mandatory) if mandatory is not None else list(MANDATORY_DEFAULT)
    if not cfg.cable.require_cable_and_pair:
        mandatory = [m for m in mandatory if m not in ("cable_id", "pair")]
    missing_cols = [m for m in mandatory if not mapping.get(m)]
    if missing_cols:
        raise InputError("Mandatory column(s) not mapped: " + ", ".join(missing_cols))
    for fld, col in mapping.items():
        if col and col not in df.columns:
            raise InputError(f"Mapped column {col!r} for {fld} does not exist in the sheet.")

    def cell(row, fld):
        col = mapping.get(fld)
        return _clean(row[col]) if col else None

    for i, row in df.iterrows():
        pair, perr = _parse_pair(row[mapping["pair"]]) if mapping.get("pair") else (None, None)
        cp, _ = _parse_pair(row[mapping["cable_pairs"]]) if mapping.get("cable_pairs") else (None, None)
        s = Signal(sid=len(signals), source_row=int(row["_source_row"]),
                   tag=cell(row, "tag"), io_raw=cell(row, "io_type"),
                   subtype_raw=cell(row, "subtype"), cable_id=cell(row, "cable_id"),
                   pair=pair, cable_pairs=cp, group=cell(row, "group"),
                   seg=(cell(row, "seg") or "").upper() or None, source_file=source_file)
        mapped_cols = {c for c in mapping.values() if c}
        s.extra = {c: row[c] for c in df.columns if c not in mapped_cols and c != "_source_row"
                   and _clean(row[c]) is not None}
        s.io_key = resolve_io_key(s.io_raw, s.subtype_raw, cfg)

        def add(code, sev, msg, needed="", **kw):
            s.findings.append(Finding(code=code, severity=sev, message=msg, sids=[s.sid],
                                      cable=s.cable_id, group=s.group, needed=needed, **kw))

        for fld in mandatory:
            val = {"tag": s.tag, "io_type": s.io_raw, "cable_id": s.cable_id, "pair": s.pair}.get(fld)
            if val is None and not (fld == "pair" and perr):
                add("MISSING_FIELD", "ERROR",
                    f"Mandatory field '{fld}' is empty (sheet row {s.source_row}).",
                    needed=f"Provide the value for '{fld}'.", data={"field": fld})
        if perr:
            add("INVALID_PAIR", "ERROR", f"{perr} (sheet row {s.source_row}).",
                needed="Provide a numeric pair number.")
        if s.io_raw and s.io_key is None:
            add("UNKNOWN_IO_TYPE", "ERROR",
                f"IO type {s.io_raw!r} is not recognised by the configured aliases.",
                needed="Confirm the IO type or add an alias in the rule profile.",
                data={"io_type": s.io_raw})
        elif s.io_key and s.io_key not in cfg.compat:
            add("UNKNOWN_COMPAT", "ERROR",
                f"No CHARM / terminal-block mapping is configured for '{s.io_key}'.",
                needed="Supply the CHARM and terminal block for this signal type.",
                rules=["R-CHARM"], data={"io_key": s.io_key})
        elif s.io_key and cfg.compat[s.io_key].status == "provisional":
            add("PROVISIONAL_COMPAT", "INFO",
                f"CHARM mapping for '{s.io_key}' is provisional and not yet validated.")
        if s.seg and not any(s.seg in p for p in cfg.segregation.conflict_pairs):
            add("UNRECOGNIZED_SEGREGATION", "WARNING",
                f"Segregation designation {s.seg!r} is not part of any configured conflict pair; "
                "it will not be segregated.", needed="Confirm the designation.")
        signals.append(s)


def _cross_checks(signals: list[Signal], cfg: RuleConfig) -> None:
    # duplicate tags -------------------------------------------------------
    by_tag: dict[str, list[Signal]] = {}
    for s in signals:
        if s.tag:
            by_tag.setdefault(s.tag_key, []).append(s)
    for key, rows in by_tag.items():
        if len(rows) < 2:
            continue
        sig = {(r.io_raw, r.cable_id, r.pair, r.group, r.seg) for r in rows}
        kind = "exact duplicate" if len(sig) == 1 else "contradictory duplicate"
        targets = rows if cfg.duplicate_policy == "exclude_all" else rows[1:]
        for r in targets:
            r.findings.append(Finding(
                code="DUPLICATE_TAG", severity="ERROR", sids=[r.sid], cable=r.cable_id, group=r.group,
                message=f"Tag {r.tag!r} appears {len(rows)} times ({kind}; sheet rows "
                        f"{', '.join(str(x.source_row) for x in rows)}).",
                needed="Remove the duplicate or correct the tag; confirm which record is valid.",
                data={"rows": [x.source_row for x in rows], "kind": kind}))
        if cfg.duplicate_policy == "keep_first":
            rows[0].findings.append(Finding(
                code="DUPLICATE_TAG_KEPT", severity="WARNING", sids=[rows[0].sid],
                message=f"Tag {rows[0].tag!r} is duplicated; first occurrence kept.",
                needed="Confirm the kept record."))
    # cable-level ----------------------------------------------------------
    by_cable: dict[str, list[Signal]] = {}
    for s in signals:
        if s.cable_key:
            by_cable.setdefault(s.cable_key, []).append(s)
    sizes = derive_cable_sizes(signals, cfg)
    for key, rows in by_cable.items():
        labels = {r.cable_id for r in rows}
        if len(labels) > 1:
            for r in rows:
                r.findings.append(Finding(
                    code="CABLE_ID_CASE", severity="WARNING", sids=[r.sid], cable=r.cable_id,
                    message=f"Cable ID appears with inconsistent capitalisation: {sorted(labels)}.",
                    needed="Confirm cable identifier spelling."))
        decl = {r.cable_pairs for r in rows if r.cable_pairs}
        if len(decl) > 1:
            for r in rows:
                r.findings.append(Finding(
                    code="CABLE_SIZE_CONFLICT", severity="ERROR", sids=[r.sid], cable=r.cable_id,
                    message=f"Cable {r.cable_id} declares conflicting sizes {sorted(decl)}.",
                    needed="Confirm the number of pairs in this cable."))
        size = sizes.get(key)
        pairs: dict[int, list[Signal]] = {}
        for r in rows:
            if r.pair is not None:
                pairs.setdefault(r.pair, []).append(r)
        for p, rs in pairs.items():
            if len(rs) > 1:
                for r in rs:
                    r.findings.append(Finding(
                        code="DUPLICATE_PAIR", severity="ERROR", sids=[r.sid], cable=r.cable_id,
                        message=f"Cable {r.cable_id} pair {p} is used by multiple signals: "
                                f"{', '.join(x.tag or '?' for x in rs)}.",
                        needed="Correct the pair numbers."))
            if size is None or p > size:
                for r in rs:
                    r.findings.append(Finding(
                        code="PAIR_OUT_OF_RANGE", severity="ERROR", sids=[r.sid], cable=r.cable_id,
                        message=f"Pair {p} exceeds the cable size "
                                f"({size if size else 'no standard size >= ' + str(p)}).",
                        needed="Confirm cable size and pair number."))


def _report(signals: list[Signal], df=None) -> ValidationReport:
    rep = ValidationReport(total_rows=len(signals))
    recs = []
    for s in signals:
        for f in s.findings:
            recs.append({"File": s.source_file, "Sheet row": s.source_row, "Tag": s.tag, "Cable": s.cable_id,
                         "Severity": f.severity, "Code": f.code, "Message": f.message})
    rep.issues = pd.DataFrame(recs, columns=["File", "Sheet row", "Tag", "Cable", "Severity", "Code", "Message"])
    for s in signals:
        codes = {f.code for f in s.findings if f.severity == "ERROR"}
        if codes:
            rep.invalid_records += 1
        if "DUPLICATE_TAG" in codes:
            rep.duplicate_tags += 1
        if "MISSING_FIELD" in codes or "INVALID_PAIR" in codes:
            rep.missing_required += 1
        if "UNKNOWN_IO_TYPE" in codes:
            rep.unrecognized_io += 1
        if any(f.code == "MISSING_FIELD" and f.data.get("field") in ("cable_id", "pair")
               for f in s.findings) or "INVALID_PAIR" in codes:
            rep.missing_cable_or_pair += 1
    rep.valid_signals = rep.total_rows - rep.invalid_records
    rep.warnings = sum(1 for s in signals for f in s.findings if f.severity == "WARNING")
    return rep


def read_master(f, filename: str, cfg: RuleConfig, sheet: Optional[str] = None,
                mapping: Optional[dict[str, Optional[str]]] = None,
                header_row: Optional[int] = None, revision_label: str = "") -> ReadResult:
    """One-call convenience: detect (or use) mapping and build signals."""
    data = _to_bytes(f)
    names = list_sheets(data, filename, cfg.max_upload_mb)
    sheet = sheet or names[0]
    df, hr = load_table(data, filename, sheet, header_row, cfg.max_upload_mb)
    if mapping is None:
        mapping, _ = detect_mapping([c for c in df.columns if c != "_source_row"])
    signals, report = build_signals(df, mapping, cfg, source_file=filename)
    return ReadResult(filename=filename, sha256=hashlib.sha256(data).hexdigest(), sheet=sheet,
                      header_row=hr, mapping=mapping, raw_df=df.copy(), signals=signals,
                      report=report, revision_label=revision_label or filename)
