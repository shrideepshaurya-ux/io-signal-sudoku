"""SYNTHETIC IO MDB generators. Nothing here is real hackathon or plant data.

``python -m src.sample_data`` writes demo workbooks into ./samples.
"""
from __future__ import annotations

import io
import random
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook

COLUMNS = ["Signal Tag", "IO Type", "Cable ID", "Pair", "Logical Group", "Segregation",
           "Characteristic", "Description"]


def _rows_for_cable(cable: str, size: int, used: int, io: str, group: str = "", seg: str = "",
                    prefix: str = "", sub: str = ""):
    out = []
    for p in range(1, used + 1):
        out.append({"Signal Tag": f"{prefix or cable}-{p:02d}", "IO Type": io, "Cable ID": cable,
                    "Pair": p, "Logical Group": group, "Segregation": seg, "Characteristic": sub,
                    "Description": f"SYNTHETIC {io} signal"})
    return out


def synthetic_basic() -> pd.DataFrame:
    """Small, fully valid dataset: all rules satisfiable."""
    rows = []
    rows += _rows_for_cable("CAB-001", 24, 22, "AI", group="G-UNIT1", seg="A", prefix="PT101")
    rows += _rows_for_cable("CAB-002", 12, 12, "DI", group="G-UNIT1", seg="B", prefix="ZS201")
    rows += _rows_for_cable("CAB-003", 6, 5, "AO", group="G-UNIT2", seg="A", prefix="FV301")
    rows += _rows_for_cable("CAB-004", 12, 10, "DO", group="G-UNIT2", seg="B", prefix="XV401")
    rows += _rows_for_cable("CAB-005", 2, 2, "DO-HC", group="G-UNIT2", prefix="SOL501")
    rows += _rows_for_cable("CAB-006", 6, 6, "DO-RELAY", group="", prefix="RLY601")
    rows += _rows_for_cable("CAB-007", 12, 9, "AI-HART", group="G-UNIT3", seg="A", prefix="FT701")
    rows += _rows_for_cable("CAB-008", 6, 4, "RTD", group="G-UNIT3", prefix="TT801")
    return pd.DataFrame(rows, columns=COLUMNS)


def synthetic_with_problems() -> pd.DataFrame:
    """Valid base + deliberate data-quality and rule problems (for TQ demos)."""
    df = synthetic_basic()
    extra = []
    extra += _rows_for_cable("CAB-101", 6, 3, "DI", group="G-X", prefix="LS101")
    # missing mandatory fields
    extra.append({"Signal Tag": "NOCABLE-01", "IO Type": "AI", "Cable ID": "", "Pair": 1,
                  "Logical Group": "", "Segregation": "", "Characteristic": "", "Description": "SYNTHETIC missing cable"})
    extra.append({"Signal Tag": "NOPAIR-01", "IO Type": "AI", "Cable ID": "CAB-102", "Pair": "",
                  "Logical Group": "", "Segregation": "", "Characteristic": "", "Description": "SYNTHETIC missing pair"})
    # unknown IO type
    extra.append({"Signal Tag": "UNK-01", "IO Type": "PULSE-TRAIN", "Cable ID": "CAB-103", "Pair": 1,
                  "Logical Group": "", "Segregation": "", "Characteristic": "", "Description": "SYNTHETIC unknown type"})
    # duplicate tags (contradictory)
    extra.append({"Signal Tag": "PT101-01", "IO Type": "DI", "Cable ID": "CAB-104", "Pair": 1,
                  "Logical Group": "", "Segregation": "", "Characteristic": "", "Description": "SYNTHETIC duplicate tag"})
    # cable carrying both A and B (impossible with continuity + segregation)
    extra.append({"Signal Tag": "MIX-01", "IO Type": "DI", "Cable ID": "CAB-105", "Pair": 1,
                  "Logical Group": "", "Segregation": "A", "Characteristic": "", "Description": "SYNTHETIC A/B conflict"})
    extra.append({"Signal Tag": "MIX-02", "IO Type": "DI", "Cable ID": "CAB-105", "Pair": 2,
                  "Logical Group": "", "Segregation": "B", "Characteristic": "", "Description": "SYNTHETIC A/B conflict"})
    # duplicate pair in one cable
    extra.append({"Signal Tag": "DUPP-01", "IO Type": "DI", "Cable ID": "CAB-106", "Pair": 1,
                  "Logical Group": "", "Segregation": "", "Characteristic": "", "Description": "SYNTHETIC dup pair"})
    extra.append({"Signal Tag": "DUPP-02", "IO Type": "DI", "Cable ID": "CAB-106", "Pair": 1,
                  "Logical Group": "", "Segregation": "", "Characteristic": "", "Description": "SYNTHETIC dup pair"})
    return pd.concat([df, pd.DataFrame(extra, columns=COLUMNS)], ignore_index=True)


def synthetic_revision_b(base: Optional[pd.DataFrame] = None) -> pd.DataFrame:
    """Rev B of ``synthetic_basic``: deletions, additions, modifications."""
    df = (base if base is not None else synthetic_basic()).copy()
    df = df[~df["Signal Tag"].isin(["FV301-05", "XV401-10"])]                 # deleted
    df.loc[df["Signal Tag"] == "PT101-03", "Logical Group"] = "G-UNIT2"        # modified attribute
    df.loc[df["Signal Tag"] == "ZS201-02", "IO Type"] = "AI"                   # modified IO type
    add = _rows_for_cable("CAB-009", 12, 8, "DI", group="G-UNIT1", seg="B", prefix="ZS901")
    add += _rows_for_cable("CAB-010", 2, 2, "AO", prefix="FV902")
    return pd.concat([df, pd.DataFrame(add, columns=COLUMNS)], ignore_index=True)


def synthetic_large(n_signals: int = 1000, seed: int = 7, group_fraction: float = 0.5) -> pd.DataFrame:
    """Random plant-like dataset used for benchmarks (valid by construction)."""
    rnd = random.Random(seed)
    rows, cab, grp_n = [], 0, 0
    types = ["AI", "AI", "AI-HART", "AO", "DI", "DI", "DO", "DO-HC", "DO-RELAY", "RTD"]
    group = ""
    left_in_group = 0
    while len(rows) < n_signals:
        cab += 1
        size = rnd.choice([24, 12, 12, 6, 6, 6, 2])
        used = max(1, int(size * rnd.uniform(0.6, 1.0)))
        used = min(used, n_signals - len(rows))
        if left_in_group <= 0:
            if rnd.random() < group_fraction:
                grp_n += 1
                group = f"G{grp_n:04d}"
                left_in_group = rnd.randint(24, 80)
            else:
                group = ""
        io = rnd.choice(types)
        seg = rnd.choice(["", "", "A", "B"])
        rows += _rows_for_cable(f"C{cab:05d}", size, used, io, group, seg, prefix=f"S{cab:05d}")
        left_in_group -= used
    return pd.DataFrame(rows[:n_signals], columns=COLUMNS)


def to_workbook_bytes(df: pd.DataFrame, sheet: str = "IO_MDB", title_rows: int = 0) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for i in range(title_rows):
        ws.append([f"SYNTHETIC DATA - NOT REAL PROJECT DATA" if i == 0 else None])
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) or v == "" else v for v in r])
        for c in ws[ws.max_row]:          # keep '=...' strings as text, never formulas
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_samples(out_dir: str = "samples") -> list[str]:
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)
    files = {
        "synthetic_io_mdb_basic.xlsx": synthetic_basic(),
        "synthetic_io_mdb_problems.xlsx": synthetic_with_problems(),
        "synthetic_io_mdb_rev_A.xlsx": synthetic_basic(),
        "synthetic_io_mdb_rev_B.xlsx": synthetic_revision_b(),
        "synthetic_io_mdb_1000.xlsx": synthetic_large(1000),
    }
    paths = []
    for name, df in files.items():
        p = d / name
        p.write_bytes(to_workbook_bytes(df, title_rows=1))
        paths.append(str(p))
    return paths


if __name__ == "__main__":
    for p in write_samples():
        print("wrote", p)
